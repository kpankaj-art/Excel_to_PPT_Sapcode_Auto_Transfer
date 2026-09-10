import streamlit as st
import pandas as pd
import re
import io
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.enum.text import MSO_VERTICAL_ANCHOR
from difflib import SequenceMatcher


# =========================================================
# PAGE SETTINGS
# =========================================================

st.set_page_config(
    page_title="SAP / Brand Transfer Tool",
    page_icon="📊",
    layout="wide"
)

st.title("📊 Transfer Workspace")
st.caption("Transfer SAP Code and/or Brand from Excel to an existing PowerPoint template using safe matching.")

# =========================================================
# PROFESSIONAL UI STYLING
# =========================================================

st.markdown(
    """
    <style>
        /* ---------- RESPONSIVE PAGE ---------- */

        .block-container {
            width: 100% !important;
            max-width: 100% !important;
            padding-top: 1rem !important;
            padding-right: 1.25rem !important;
            padding-bottom: 2rem !important;
            padding-left: 1.25rem !important;
            overflow-x: hidden !important;
        }

        .main .block-container {
            overflow-x: hidden !important;
        }

        /* Keep the sidebar compact so it does not push the workspace
           outside the browser viewport. */
        [data-testid="stSidebar"] {
            width: 290px !important;
            min-width: 290px !important;
            max-width: 290px !important;
        }

        [data-testid="stSidebar"] > div:first-child {
            width: 290px !important;
        }

        [data-testid="stSidebar"] .block-container {
            width: 100% !important;
            padding: 1rem 0.85rem 1.5rem 0.85rem !important;
            overflow-x: hidden !important;
        }

        /* ---------- MAIN TITLE ---------- */

        h1 {
            font-size: clamp(1.65rem, 2.2vw, 2.35rem) !important;
            line-height: 1.15 !important;
            margin-top: 0 !important;
            margin-bottom: 0.35rem !important;
            white-space: normal !important;
            overflow-wrap: anywhere !important;
        }

        .main p {
            max-width: 100% !important;
            overflow-wrap: anywhere !important;
        }

        /* ---------- SIDEBAR ---------- */

        [data-testid="stSidebar"] h2,
        [data-testid="stSidebar"] h3 {
            font-size: 1.05rem !important;
            line-height: 1.25 !important;
        }

        [data-testid="stSidebar"] .stCaption {
            font-size: 0.78rem !important;
        }

        [data-testid="stSidebar"] [data-testid="stFileUploader"] {
            margin-bottom: 0.45rem !important;
        }

        /* Make uploader boxes shorter and fit narrow screens. */
        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] {
            min-height: 92px !important;
            padding: 0.65rem !important;
        }

        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzone"] > div {
            gap: 0.35rem !important;
        }

        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] {
            padding: 0 !important;
        }

        [data-testid="stSidebar"] [data-testid="stFileUploaderDropzoneInstructions"] > div {
            font-size: 0.78rem !important;
        }

        [data-testid="stSidebar"] .stButton button {
            min-height: 44px !important;
            width: 100% !important;
            white-space: normal !important;
            line-height: 1.2 !important;
        }

        /* ---------- TABLES ---------- */

        [data-testid="stDataFrame"] {
            width: 100% !important;
            max-width: 100% !important;
        }

        /* ---------- MOBILE / SMALL WINDOW ---------- */

        @media (max-width: 900px) {

            .block-container {
                padding-left: 0.75rem !important;
                padding-right: 0.75rem !important;
            }

            [data-testid="stSidebar"] {
                width: 260px !important;
                min-width: 260px !important;
                max-width: 260px !important;
            }

            [data-testid="stSidebar"] > div:first-child {
                width: 260px !important;
            }

            h1 {
                font-size: 1.55rem !important;
            }
        }

        @media (max-width: 640px) {

            .block-container {
                padding-left: 0.5rem !important;
                padding-right: 0.5rem !important;
            }

            h1 {
                font-size: 1.35rem !important;
            }
        }

        /* ---------- GENERAL ---------- */

        .app-card {
            border: 1px solid rgba(128,128,128,0.25);
            border-radius: 12px;
            padding: 16px;
            margin-bottom: 14px;
        }

        .section-title {
            font-size: 1.05rem;
            font-weight: 750;
            margin-top: 8px;
            margin-bottom: 10px;
        }

        /* Prevent long filenames and messages from causing horizontal overflow. */
        [data-testid="stSidebar"] * {
            overflow-wrap: anywhere;
            word-break: break-word;
        }

        .main * {
            max-width: 100%;
        }
    </style>
    """,
    unsafe_allow_html=True
)


# =========================================================
# HELPER FUNCTIONS
# =========================================================

def safe_filename(name):
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip()


def normalize_text(value):
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except:
        pass

    value = str(value).strip().upper()
    value = value.replace("\n", " ").replace("\r", " ")
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"[^A-Z0-9 ]", " ", value)
    value = re.sub(r"\s+", " ", value).strip()

    return value


def normalize_phone(value):
    if value is None:
        return []

    try:
        if pd.isna(value):
            return []
    except:
        pass

    value = str(value).strip()

    if "E+" in value.upper():
        try:
            value = str(int(float(value)))
        except:
            pass

    if value.endswith(".0"):
        value = value[:-2]

    numbers = re.findall(r"\d{10,}", value)

    if not numbers:
        digits = re.sub(r"\D", "", value)
        if len(digits) >= 10:
            numbers = [digits]

    return list(dict.fromkeys(numbers))


def normalize_size(value):
    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except:
        pass

    value = str(value).upper().strip()
    value = value.replace(" ", "")
    value = value.replace("*", "X")
    value = value.replace("×", "X")

    match = re.search(r"(\d+(?:\.\d+)?)X(\d+(?:\.\d+)?)", value)

    if match:
        return f"{match.group(1)}X{match.group(2)}"

    return value


def sizes_equal(size1, size2):
    s1 = normalize_size(size1)
    s2 = normalize_size(size2)
    return s1 != "" and s2 != "" and s1 == s2


def names_similar(name1, name2, threshold=0.88):
    n1 = normalize_text(name1)
    n2 = normalize_text(name2)

    if not n1 or not n2:
        return False

    if n1 == n2:
        return True

    return SequenceMatcher(None, n1, n2).ratio() >= threshold


# =========================================================
# EXCEL COLUMN DETECTION
# =========================================================

def clean_column_name(col):
    return normalize_text(col)


def detect_columns(df):

    columns = list(df.columns)

    normalized_columns = {
        col: clean_column_name(col)
        for col in columns
    }

    mapping = {}

    name_aliases = [
        "DEALER / NAME",
        "DEALER NAME",
        "OUTLET NAME",
        "OUTLET",
        "DEALER",
        "CUSTOMER NAME",
        "NAME"
    ]

    contact_aliases = [
        "DEALER / CONTACT",
        "DEALER CONTACT",
        "CONTACT NO",
        "CONTACT NUMBER",
        "CONTACT",
        "PHONE",
        "MOBILE",
        "MOBILE NO"
    ]

    sap_aliases = [
        "SAPCODE",
        "SAP CODE",
        "DEALER CODE",
        "CUSTOMER CODE",
        "CUSTOMERCODE",
        "CUSTOMER ID",
        "SAP"
    ]

    address_aliases = [
        "DEALER / ADDERSS",
        "DEALER / ADDRESS",
        "DEALER ADDRESS",
        "ADDRESS"
    ]

    district_aliases = [
        "DISTRICT NAME",
        "DISTRICT"
    ]

    type_aliases = [
        "TYPE",
        "MEDIA TYPE",
        "MEDIA"
    ]

    width_aliases = [
        "W",
        "WIDTH"
    ]

    height_aliases = [
        "H",
        "HEIGHT"
    ]

    size_aliases = [
        "SIZE",
        "DIMENSION",
        "DIMENSIONS"
    ]

    brand_aliases = [
        "BRAND",
        "BRAND NAME",
        "BRAND_NAME",
        "BRANDNAME"
    ]

    def find_column(aliases, exclude=None):
        exclude = exclude or []
        alias_norms = [clean_column_name(x) for x in aliases]

        # Exact match first
        for col, norm in normalized_columns.items():
            if col in exclude:
                continue
            if norm in alias_norms:
                return col

        # Partial match second
        for col, norm in normalized_columns.items():
            if col in exclude:
                continue
            for alias_norm in alias_norms:
                if alias_norm and alias_norm in norm:
                    return col

        return None

    mapping["name"] = find_column(name_aliases)

    mapping["contact"] = find_column(
        contact_aliases,
        exclude=[mapping["name"]] if mapping["name"] else []
    )

    mapping["sap"] = find_column(sap_aliases)
    mapping["address"] = find_column(address_aliases)
    mapping["district"] = find_column(district_aliases)
    mapping["type"] = find_column(type_aliases)
    mapping["width"] = find_column(width_aliases)
    mapping["height"] = find_column(height_aliases)
    mapping["brand"] = find_column(brand_aliases)

    # Give Width + Height priority for Size
    if mapping["width"] and mapping["height"]:
        mapping["size"] = None
    else:
        mapping["size"] = find_column(
            size_aliases,
            exclude=[mapping["type"]] if mapping["type"] else []
        )

    return mapping


# =========================================================
# POWERPOINT PARSING
# =========================================================

def parse_label_line(line):

    if not line:
        return None, None

    line = str(line).strip()

    patterns = [
        (r"^\s*Outlet\s*Name\s*:\s*(.*)$", "name"),
        (r"^\s*Address\s*:\s*(.*)$", "address"),
        (r"^\s*Contact\s*No\s*:\s*(.*)$", "contact"),
        (r"^\s*Contact\s*Number\s*:\s*(.*)$", "contact"),
        (r"^\s*District\s*:\s*(.*)$", "district"),
        (r"^\s*Sap\s*code\s*:\s*(.*)$", "sap"),
        (r"^\s*Sapcode\s*:\s*(.*)$", "sap"),
        (r"^\s*SAP\s*Code\s*:\s*(.*)$", "sap"),
        (r"^\s*Size\s*:\s*(.*)$", "size"),
        (r"^\s*Media\s*Type\s*:\s*(.*)$", "type"),
        (r"^\s*Type\s*:\s*(.*)$", "type"),
        (r"^\s*Remarks\s*:\s*(.*)$", "remarks"),
        (r"^\s*Qty\s*:\s*(.*)$", "qty"),
    ]

    for pattern, field in patterns:
        match = re.match(pattern, line, flags=re.IGNORECASE)
        if match:
            return field, match.group(1).strip()

    return None, None


def extract_shape_text(shape):
    try:
        if hasattr(shape, "text"):
            return shape.text or ""
    except:
        pass
    return ""


def extract_ppt_fields(prs):

    slides_data = []

    for slide_number, slide in enumerate(prs.slides, start=1):

        data = {
            "slide": slide_number,
            "name": "",
            "address": "",
            "contact": "",
            "district": "",
            "sap": "",
            "size": "",
            "type": "",
            "remarks": "",
            "qty": "",
            "info_shape_index": None
        }

        for shape_index, shape in enumerate(slide.shapes):

            text = extract_shape_text(shape)

            if not text:
                continue

            for line in text.splitlines():

                field, value = parse_label_line(line)

                if field:
                    data[field] = value

            lower_text = text.lower()

            if (
                "outlet name" in lower_text
                and "contact" in lower_text
                and "district" in lower_text
            ):
                data["info_shape_index"] = shape_index

        slides_data.append(data)

    return slides_data


# =========================================================
# SIZE FROM EXCEL
# =========================================================

def get_excel_size(row, mapping):

    width_col = mapping.get("width")
    height_col = mapping.get("height")
    size_col = mapping.get("size")

    if width_col and height_col:

        width = row.get(width_col, "")
        height = row.get(height_col, "")

        width_ok = str(width).strip() != "" and str(width).lower() != "nan"
        height_ok = str(height).strip() != "" and str(height).lower() != "nan"

        if width_ok and height_ok:
            return normalize_size(f"{width}X{height}")

    if size_col:
        return normalize_size(row.get(size_col, ""))

    return ""


# =========================================================
# SAFE MATCHING
# =========================================================

def find_best_slide(excel_row, ppt_slides, mapping, used_slides):

    excel_name = excel_row.get(mapping.get("name"), "")
    excel_contact = excel_row.get(mapping.get("contact"), "")
    excel_size = get_excel_size(excel_row, mapping)

    excel_phones = normalize_phone(excel_contact)

    if not excel_name:
        return None, "Excel dealer/outlet name is missing"

    if not excel_phones:
        return None, "Excel contact number is missing or invalid"

    candidates = []

    for ppt in ppt_slides:

        slide_no = ppt["slide"]

        # Each PowerPoint slide can be used only once
        if slide_no in used_slides:
            continue

        if not names_similar(excel_name, ppt.get("name", "")):
            continue

        ppt_phones = normalize_phone(ppt.get("contact", ""))

        if not ppt_phones:
            continue

        contact_match = any(
            phone in ppt_phones
            for phone in excel_phones
        )

        if not contact_match:
            continue

        candidates.append(ppt)

    if not candidates:
        return None, "No Name + Contact match found"

    # Unique Name + Contact match
    if len(candidates) == 1:
        return candidates[0], "Matched by Name + Contact"

    # Duplicate Name + Contact requires Size verification
    if not excel_size:
        return None, "Multiple Name + Contact matches found; Size is required for verification"

    size_candidates = []

    for ppt in candidates:

        if sizes_equal(
            excel_size,
            ppt.get("size", "")
        ):
            size_candidates.append(ppt)

    if len(size_candidates) == 1:
        return size_candidates[0], "Matched by Name + Contact + Size"

    if len(size_candidates) > 1:
        return None, "Multiple Name + Contact + Size matches found; transfer skipped"

    return None, "Multiple Name + Contact matches found, but Size did not match"


# =========================================================
# TEXT FORMATTING
# =========================================================

def clear_and_add_bold_text(shape, lines, font_size=15):

    shape.text_frame.clear()

    for i, line in enumerate(lines):

        if i == 0:
            paragraph = shape.text_frame.paragraphs[0]
        else:
            paragraph = shape.text_frame.add_paragraph()

        paragraph.text = ""

        run = paragraph.add_run()
        run.text = str(line)
        run.font.size = Pt(font_size)
        run.font.bold = True

        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)

    try:
        shape.text_frame.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    except:
        pass


# =========================================================
# UPDATE EXISTING INFORMATION BOX
# =========================================================

def add_sap_to_info_shape(shape, sap_code):

    old_text = shape.text or ""
    fields = []

    for line in old_text.splitlines():

        field, value = parse_label_line(line)

        if field in [
            "name",
            "address",
            "contact",
            "district"
        ]:
            fields.append((field, value))

    final_lines = []
    inserted = False

    label_map = {
        "name": "Outlet Name",
        "address": "Address",
        "contact": "Contact No",
        "district": "District"
    }

    for field, value in fields:

        final_lines.append(
            f"{label_map[field]} : {value}"
        )

        if field == "district":

            final_lines.append(
                f"Sapcode     : {sap_code}"
            )

            inserted = True

    if not inserted:
        final_lines.append(
            f"Sapcode     : {sap_code}"
        )

    # Reduce font size when the information block contains many lines
    if len(final_lines) >= 8:
        font_size = 12
    elif len(final_lines) >= 7:
        font_size = 13
    else:
        font_size = 15

    clear_and_add_bold_text(
        shape,
        final_lines,
        font_size
    )


# =========================================================
# BRAND TEXTBOX
# =========================================================

def find_info_shape(slide, info_shape_index):

    if info_shape_index is not None:

        try:
            if info_shape_index < len(slide.shapes):
                shape = slide.shapes[info_shape_index]

                if (
                    "outlet name" in extract_shape_text(shape).lower()
                    and "contact" in extract_shape_text(shape).lower()
                ):
                    return shape
        except:
            pass

    for shape in slide.shapes:

        text = extract_shape_text(shape).lower()

        if (
            "outlet name" in text
            and "contact" in text
            and "district" in text
        ):
            return shape

    return None


def add_brand_in_green_area(slide, info_shape, brand_text):

    if not brand_text:
        return False

    # -----------------------------------------------------
    # BRAND AREA:
    # Position the brand within the right-side area of the information box.
    # Coordinates are calculated relative to the existing information box.
    # -----------------------------------------------------

    left = info_shape.left
    top = info_shape.top
    width = info_shape.width
    height = info_shape.height

    # Use approximately the right-side 35% of the information box.
    brand_left = left + int(width * 0.66)
    brand_top = top + int(height * 0.08)
    brand_width = int(width * 0.31)
    brand_height = int(height * 0.84)

    # Create a transparent text box
    textbox = slide.shapes.add_textbox(
        brand_left,
        brand_top,
        brand_width,
        brand_height
    )

    textbox.name = "AUTO_BRAND"

    # No fill / no border.
    try:
        textbox.fill.background()
    except:
        pass

    try:
        textbox.line.fill.background()
    except:
        pass

    tf = textbox.text_frame
    tf.clear()

    # Automatically adjust font size for long brand names
    brand_length = len(str(brand_text))

    if brand_length <= 18:
        font_size = 18
    elif brand_length <= 28:
        font_size = 16
    elif brand_length <= 40:
        font_size = 14
    else:
        font_size = 12

    paragraph = tf.paragraphs[0]
    paragraph.text = ""
    paragraph.alignment = 1

    run = paragraph.add_run()
    run.text = str(brand_text).strip()
    run.font.bold = True
    run.font.size = Pt(font_size)

    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)

    try:
        tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    except:
        pass

    return True


# =========================================================
# REMOVE PREVIOUSLY GENERATED BRAND BOXES
# =========================================================

def remove_old_auto_brand(slide):

    # XML se AUTO_BRAND textbox remove
    shapes_to_remove = []

    for shape in slide.shapes:

        try:
            if shape.name == "AUTO_BRAND":
                shapes_to_remove.append(shape)
        except:
            pass

    for shape in shapes_to_remove:

        try:
            sp = shape._element
            sp.getparent().remove(sp)
        except:
            pass


# =========================================================
# UPDATE POWERPOINT
# =========================================================

def update_ppt(prs, matching_results, add_mode):

    updated_count = 0
    failed_count = 0

    add_sap = add_mode in [
        "SAP Code",
        "Both (SAP Code + Brand)"
    ]

    add_brand = add_mode in [
        "Brand",
        "Both (SAP Code + Brand)"
    ]

    for result in matching_results:

        if not result["matched"]:
            continue

        slide_number = result["slide"]
        slide = prs.slides[slide_number - 1]

        target_shape = find_info_shape(
            slide,
            result.get("info_shape_index")
        )

        if target_shape is None:

            failed_count += 1
            continue

        # Remove previously generated Brand box
        remove_old_auto_brand(slide)

        # -------------------------------------------------
        # SAP CODE
        # -------------------------------------------------

        if add_sap:

            add_sap_to_info_shape(
                target_shape,
                result["sap"]
            )

        # -------------------------------------------------
        # ADD BRAND
        # -------------------------------------------------

        if add_brand:

            brand_value = result.get("brand", "")

            if brand_value:

                add_brand_in_green_area(
                    slide,
                    target_shape,
                    brand_value
                )

        updated_count += 1

    return updated_count, failed_count


# =========================================================
# MATCHING REPORT
# =========================================================

def create_report_excel(results):

    report_rows = []

    for result in results:

        report_rows.append({

            "Excel Row":
                result.get("excel_row", ""),

            "Excel Name":
                result.get("excel_name", ""),

            "Excel Contact":
                result.get("excel_contact", ""),

            "Excel Size":
                result.get("excel_size", ""),

            "SAP Code":
                result.get("sap", ""),

            "Brand":
                result.get("brand", ""),

            "PPT Slide":
                result.get("slide", ""),

            "PPT Name":
                result.get("ppt_name", ""),

            "PPT Contact":
                result.get("ppt_contact", ""),

            "PPT Size":
                result.get("ppt_size", ""),

            "Status":
                "MATCHED"
                if result.get("matched")
                else "NOT MATCHED",

            "Reason":
                result.get("reason", "")
        })

    df_report = pd.DataFrame(report_rows)

    output = io.BytesIO()

    with pd.ExcelWriter(
        output,
        engine="openpyxl"
    ) as writer:

        df_report.to_excel(
            writer,
            index=False,
            sheet_name="Matching Report"
        )

    output.seek(0)

    return output.getvalue()


# =========================================================
# LEFT SIDEBAR - FILES AND ACTIONS
# =========================================================

with st.sidebar:

    st.markdown("## ⚙️ Actions")
    st.caption("Upload your files and choose the transfer operation.")

    st.markdown("### 1. Upload Files")

    excel_file = st.file_uploader(
        "Upload Excel File",
        type=["xlsx", "xls"],
        key="excel_upload"
    )

    ppt_file = st.file_uploader(
        "Upload PowerPoint Template",
        type=["pptx"],
        key="ppt_upload"
    )

    st.markdown("---")

    st.markdown("### 2. Transfer Content")

    add_mode = st.selectbox(
        "Select What to Add",
        [
            "SAP Code",
            "Brand",
            "Both (SAP Code + Brand)"
        ],
        key="add_mode"
    )

    st.caption(
        "SAP Code is added below District. "
        "Brand is placed in the designated brand area."
    )

    st.markdown("---")

    ready = excel_file is not None and ppt_file is not None

    if ready:
        st.success("Files are ready.")
    else:
        st.info("Upload both files to continue.")

    process_button = st.button(
        "🚀 Transfer / Update PowerPoint",
        type="primary",
        use_container_width=True,
        disabled=not ready
    )

    if ready:
        st.markdown("---")
        st.caption("Selected files")
        st.write(f"**Excel:** {excel_file.name}")
        st.write(f"**PowerPoint:** {ppt_file.name}")


# =========================================================
# MAIN PROCESS
# =========================================================



if excel_file and ppt_file:

    st.markdown(
        '<div class="section-title">PowerPoint Transfer Preview</div>',
        unsafe_allow_html=True
    )

    st.divider()

    # -----------------------------------------------------
    # CONTENT SELECTION
    # -----------------------------------------------------

    st.subheader("2️⃣ Select What to Add")

    add_mode = st.selectbox(
        "Select Option",
        [
            "SAP Code",
            "Brand",
            "Both (SAP Code + Brand)"
        ]
    )

    st.info(
        "SAP Code will be added below District in the information box. "
        "Brand will be added as bold text in the designated brand area."
    )

    if process_button:

        try:

            # -------------------------------------------------
            # Read Excel
            # -------------------------------------------------

            df = pd.read_excel(
                excel_file
            )

            # -------------------------------------------------
            # Detect columns
            # -------------------------------------------------

            mapping = detect_columns(df)

            st.subheader("3️⃣ Detected Excel Columns")

            display_mapping = {

                "Outlet / Dealer Name":
                    mapping.get("name"),

                "Contact":
                    mapping.get("contact"),

                "SAP Code":
                    mapping.get("sap"),

                "Brand":
                    mapping.get("brand"),

                "Address":
                    mapping.get("address"),

                "District":
                    mapping.get("district"),

                "Media Type":
                    mapping.get("type"),

                "Width":
                    mapping.get("width"),

                "Height":
                    mapping.get("height"),

                "Size":
                    "W + H"
                    if mapping.get("width")
                    and mapping.get("height")
                    else mapping.get("size")
            }

            mapping_df = pd.DataFrame(
                list(display_mapping.items()),
                columns=[
                    "Field",
                    "Excel Column"
                ]
            )

            st.dataframe(
                mapping_df,
                use_container_width=True,
                hide_index=True
            )

            # -------------------------------------------------
            # REQUIRED MATCHING COLUMNS
            # -------------------------------------------------

            missing = []

            if not mapping.get("name"):
                missing.append("Dealer / Outlet Name")

            if not mapping.get("contact"):
                missing.append("Contact")

            # SAP CODE is required only when SAP Code is selected
            if add_mode in [
                "SAP Code",
                "Both (SAP Code + Brand)"
            ]:
                if not mapping.get("sap"):
                    missing.append("SAP Code / Customer Code")

            # ADD BRAND is required only when Brand is selected
            if add_mode in [
                "Brand",
                "Both (SAP Code + Brand)"
            ]:
                if not mapping.get("brand"):
                    missing.append("Brand")

            if missing:

                st.error(
                    "The following required Excel columns could not be detected: "
                    + ", ".join(missing)
                )

                st.stop()

            # -------------------------------------------------
            # Load PPT
            # -------------------------------------------------

            ppt_bytes = ppt_file.getvalue()

            prs = Presentation(
                io.BytesIO(ppt_bytes)
            )

            # -------------------------------------------------
            # Extract PPT
            # -------------------------------------------------

            ppt_slides = extract_ppt_fields(prs)

            st.subheader("4️⃣ Detected PowerPoint Data")

            ppt_preview = []

            for item in ppt_slides:

                ppt_preview.append({

                    "Slide":
                        item["slide"],

                    "Outlet Name":
                        item["name"],

                    "Contact":
                        item["contact"],

                    "Size":
                        item["size"],

                    "Media Type":
                        item["type"],

                    "District":
                        item["district"]
                })

            st.dataframe(
                pd.DataFrame(ppt_preview),
                use_container_width=True,
                hide_index=True
            )

            # -------------------------------------------------
            # MATCHING
            # -------------------------------------------------

            st.subheader("5️⃣ Safe Matching")

            used_slides = set()
            matching_results = []

            for excel_index, row in df.iterrows():

                excel_name = row.get(
                    mapping.get("name"),
                    ""
                )

                excel_contact = row.get(
                    mapping.get("contact"),
                    ""
                )

                sap_code = row.get(
                    mapping.get("sap"),
                    ""
                ) if mapping.get("sap") else ""

                brand_value = row.get(
                    mapping.get("brand"),
                    ""
                ) if mapping.get("brand") else ""

                excel_size = get_excel_size(
                    row,
                    mapping
                )

                # Clean blank Brand values.
                try:
                    if pd.isna(brand_value):
                        brand_value = ""
                except:
                    pass

                brand_value = str(
                    brand_value
                ).strip()

                matched_slide, reason = find_best_slide(
                    row,
                    ppt_slides,
                    mapping,
                    used_slides
                )

                if matched_slide is not None:

                    slide_no = matched_slide["slide"]

                    used_slides.add(slide_no)

                    matching_results.append({

                        "excel_row":
                            excel_index + 2,

                        "excel_name":
                            str(excel_name),

                        "excel_contact":
                            str(excel_contact),

                        "excel_size":
                            excel_size,

                        "sap":
                            str(sap_code),

                        "brand":
                            brand_value,

                        "slide":
                            slide_no,

                        "ppt_name":
                            matched_slide["name"],

                        "ppt_contact":
                            matched_slide["contact"],

                        "ppt_size":
                            matched_slide["size"],

                        "info_shape_index":
                            matched_slide["info_shape_index"],

                        "matched":
                            True,

                        "reason":
                            reason
                    })

                else:

                    matching_results.append({

                        "excel_row":
                            excel_index + 2,

                        "excel_name":
                            str(excel_name),

                        "excel_contact":
                            str(excel_contact),

                        "excel_size":
                            excel_size,

                        "sap":
                            str(sap_code),

                        "brand":
                            brand_value,

                        "slide":
                            "",

                        "ppt_name":
                            "",

                        "ppt_contact":
                            "",

                        "ppt_size":
                            "",

                        "info_shape_index":
                            None,

                        "matched":
                            False,

                        "reason":
                            reason
                    })

            # -------------------------------------------------
            # Summary
            # -------------------------------------------------

            matched_count = sum(
                1
                for x in matching_results
                if x["matched"]
            )

            unmatched_count = (
                len(matching_results)
                - matched_count
            )

            col1, col2, col3 = st.columns(3)

            col1.metric(
                "Excel Rows",
                len(df)
            )

            col2.metric(
                "Matched",
                matched_count
            )

            col3.metric(
                "Not Matched",
                unmatched_count
            )

            # -------------------------------------------------
            # RESULT table
            # -------------------------------------------------

            result_preview = []

            for result in matching_results:

                result_preview.append({

                    "Excel Row":
                        result["excel_row"],

                    "Excel Name":
                        result["excel_name"],

                    "Contact":
                        result["excel_contact"],

                    "Size":
                        result["excel_size"],

                    "SAP Code":
                        result["sap"],

                    "Brand":
                        result["brand"],

                    "PPT Slide":
                        result["slide"],

                    "Status":
                        "✅ MATCHED"
                        if result["matched"]
                        else "❌ NOT MATCHED",

                    "Reason":
                        result["reason"]
                })

            st.dataframe(
                pd.DataFrame(result_preview),
                use_container_width=True,
                hide_index=True
            )

            # -------------------------------------------------
            # TRANSFER
            # -------------------------------------------------

            st.divider()

            if st.button(
                "🚀 Transfer / Add to PPT",
                type="primary",
                use_container_width=True
            ):

                with st.spinner(
                    "Updating the PowerPoint file..."
                ):

                    updated_count, failed_count = update_ppt(
                        prs,
                        matching_results,
                        add_mode
                    )

                    # -----------------------------------------
                    # Save PPT
                    # -----------------------------------------

                    ppt_output = io.BytesIO()

                    prs.save(
                        ppt_output
                    )

                    ppt_output.seek(0)

                    final_ppt_bytes = (
                        ppt_output.getvalue()
                    )

                    # -----------------------------------------
                    # Preserve original PowerPoint filename and append _Update
                    # -----------------------------------------

                    original_ppt_name = ppt_file.name

                    base_ppt_name = re.sub(
                        r"\.pptx$",
                        "",
                        original_ppt_name,
                        flags=re.IGNORECASE
                    )

                    updated_name = safe_filename(
                        base_ppt_name + "_Update.pptx"
                    )

                    # -----------------------------------------
                    # Generate matching report filename from the Excel filename
                    # -----------------------------------------

                    original_excel_name = excel_file.name

                    excel_base_name = re.sub(
                        r"\.(xlsx|xls)$",
                        "",
                        original_excel_name,
                        flags=re.IGNORECASE
                    )

                    report_name = safe_filename(
                        excel_base_name + "_Matching_Report.xlsx"
                    )

                    # -----------------------------------------
                    # RESULT
                    # -----------------------------------------

                    st.success(
                        f"Done! {updated_count} PowerPoint slide(s) were successfully updated."
                    )

                    if failed_count > 0:

                        st.warning(
                            f"{failed_count} matched slide(s) could not be updated."
                        )

                    # -----------------------------------------
                    # Download Updated PPT
                    # -----------------------------------------

                    st.download_button(
                        label="📥 Download Updated PowerPoint",
                        data=final_ppt_bytes,
                        file_name=updated_name,
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "presentationml.presentation"
                        ),
                        use_container_width=True
                    )

                    # -----------------------------------------
                    # Download matching report.
                    # -----------------------------------------

                    report_bytes = create_report_excel(
                        matching_results
                    )

                    st.download_button(
                        label="📊 Download Matching Report",
                        data=report_bytes,
                        file_name=report_name,
                        mime=(
                            "application/vnd.openxmlformats-officedocument."
                            "spreadsheetml.sheet"
                        ),
                        use_container_width=True
                    )

                    st.info(
                        "Safety rule: A unique Name + Contact match is required. "
                        "If the same Name + Contact appears multiple times, Size is used for verification. "
                        "SAP Code and/or Brand will not be transferred when a unique match cannot be confirmed."
                    )
        except Exception as e:

            st.error("An error occurred:")

            st.exception(e)
